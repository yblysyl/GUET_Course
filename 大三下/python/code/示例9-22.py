import random
import os
#需要安装扩展库python-docx
from docx import Document

columnsNumber = 4

def main(rowsNumber=20, grade=4):
    #默认生成20行小学四年级口算题
    if grade < 3:
        operators = '＋－'
        biggest = 20
    elif grade <= 4:
        operators = '＋－×÷'
        biggest = 100
    elif grade == 5:
        operators = '＋－×÷('
        biggest = 100
        
    document = Document()
    #创建表格
    table = document.add_table(rows=rowsNumber, cols=columnsNumber)
    #遍历每个单元格
    for row in range(rowsNumber):
        for col in range(columnsNumber):
            first = random.randint(1, biggest)
            second = random.randint(1, biggest)
            operator = random.choice(operators)            
            if operator != '(':
                if operator == '－':
                    #如果是减法口算题，确保结果为正数
                    if first < second:
                        first, second = second, first
                r = str(first).ljust(2, ' ') + ' ' + operator  + str(second).ljust(2, ' ') + '='
            else:
                #生成带括号的口算题，需要3个数字和2个运算符
                third = random.randint(1, 100)
                while True:
                    o1 = random.choice(operators)
                    o2 = random.choice(operators)
                    if o1 != '(' and o2 != '(':
                        break
                rr = random.randint(1, 100)
                if rr > 50:
                    if o2 == '－':
                        if second < third:
                            second, third = third, second
                    r = str(first).ljust(2, ' ') + o1 + ' ( ' \
                       + str(second).ljust(2, ' ') + o2 + str(third).ljust(2, ' ') + ')='
                else:
                    if o1 == '-':
                        if first < second:
                            first, second = second, first
                    r = '(' + str(first).ljust(2, ' ') + o1 \
                       + str(second).ljust(2, ' ') + ')' \
                       + o2 + str(third).ljust(2, ' ') + '='
            #获取指定单元格并写入口算题
            cell = table.cell(row, col)
            cell.text = r
    document.save('kousuan.docx') 
    os.startfile('kousuan.docx')

main(grade=5)
